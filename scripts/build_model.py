"""Build model/qcom_revenue_forecast_model.xlsx from data/processed CSVs."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

ROOT = "."
OUT = f"{ROOT}/model/qcom_revenue_forecast_model.xlsx"

# ---------- data ----------
act = list(csv.DictReader(open(f"{ROOT}/data/processed/segment_actuals.csv")))
gui = list(csv.DictReader(open(f"{ROOT}/data/processed/guidance.csv")))
src = {r["source_id"]: r for r in csv.DictReader(open(f"{ROOT}/docs/sources.csv"))}

Q = [(fy, q) for fy in (2024, 2025, 2026, 2027) for q in (1, 2, 3, 4)]
N_ACT = 11                      # FY24Q1..FY26Q3
FC0 = 3                         # first quarter column (C)
col = lambda i: L(FC0 + i)      # quarter index -> column letter
lab = lambda fy, q: f"FY{str(fy)[2:]} Q{q}"

A = {}
for r in act:
    k = (int(r["fiscal_year"]), int(r["fiscal_quarter"]))
    A.setdefault(k, {"pe": r["period_end"], "src": r["source_id"]})
    A[k][(r["segment"], r["sub_segment"], "rev")] = float(r["revenue_musd"])
    if r["ebt_musd"]:
        A[k][(r["segment"], r["sub_segment"], "ebt")] = float(r["ebt_musd"])
G = {}
for r in gui:
    k = (int(r["fiscal_year"]), int(r["fiscal_quarter"]))
    G.setdefault(k, {"issued": r["issued_date"], "src": r["source_id"]})
    G[k][r["metric"]] = (float(r["low_musd"]), float(r["high_musd"]))

# ---------- styles ----------
F = "Arial"
f_base = Font(name=F, size=10)
f_in = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
f_bold = Font(name=F, size=10, bold=True)
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title = Font(name=F, size=14, bold=True, color="1F2A44")
f_sub = Font(name=F, size=10, italic=True, color="595959")
fill_hdr = PatternFill("solid", fgColor="1F2A44")
fill_sec = PatternFill("solid", fgColor="DCE3EE")
fill_key = PatternFill("solid", fgColor="FFFF00")
fill_fc = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="808080")
top_border = Border(top=thin)
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
CENTER = Alignment(horizontal="center", vertical="center")

wb = Workbook()


def setup(ws, title, subtitle, widths=None):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title; ws["A1"].font = f_title
    ws["A2"] = subtitle; ws["A2"].font = f_sub
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12
    for c, w in (widths or {}).items():
        ws.column_dimensions[c].width = w


def put(ws, ref, v, font=f_base, fmt=None, fill=None, align=None, border=None):
    c = ws[ref]; c.value = v; c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align: c.alignment = align
    if border: c.border = border
    return c


def header_row(ws, row, labels_start_col, labels, first=("Line item", "Unit")):
    put(ws, f"A{row}", first[0], f_hdr, fill=fill_hdr)
    put(ws, f"B{row}", first[1], f_hdr, fill=fill_hdr, align=CENTER)
    for j, t in enumerate(labels):
        put(ws, f"{L(labels_start_col + j)}{row}", t, f_hdr, fill=fill_hdr, align=CENTER)


def section(ws, row, text, last_col):
    for c in range(1, last_col + 1):
        ws.cell(row=row, column=c).fill = fill_sec
    put(ws, f"A{row}", text, f_bold, fill=fill_sec)


def qwidths(ws, n=16, w=11):
    for i in range(n):
        ws.column_dimensions[col(i)].width = w


# =====================================================================
# Cover
# =====================================================================
ws = wb.active; ws.title = "Cover"
setup(ws, "Qualcomm Financial Revenue Forecast & Earnings Variance Analysis",
      "Forecasting Qualcomm's shift from phones to cars & AI  |  $ in millions unless noted")
ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 90
put(ws, "A4", "Active scenario", f_bold)
put(ws, "B4", "=Assumptions!C5", f_link)
put(ws, "A5", "Model checks", f_bold)
put(ws, "B5", "=Checks!C4", f_link)

section(ws, 7, "Tab guide", 2)
tabs = [
    ("Assumptions", "Scenario selector (1=Base, 2=Bull, 3=Bear) and every forecast driver with rationale"),
    ("Actuals", "Reported segment revenue and EBT, FY24 Q1 - FY26 Q3, from earnings releases (S1-S11)"),
    ("Guidance", "Company revenue guidance ranges issued one quarter ahead, FY24 Q2 - FY26 Q4"),
    ("Variance", "Actual vs. guidance midpoint by segment: $ and % beat/miss, position in range"),
    ("Forecast", "Quarterly actuals + forecast FY26 Q4 - FY27 Q4 driven by the active scenario"),
    ("Summary", "Fiscal-year roll-up and the core question: does diversification offset handset decline?"),
    ("Scenarios", "FY26E / FY27E revenue and EBT under Base, Bull and Bear side by side"),
    ("Sensitivity", "FY27E revenue across handset growth x automotive growth"),
    ("Checks", "Integrity tests: stream tie-outs, guidance math, annual roll-ups, selector"),
    ("Tableau_Revenue", "Long-format revenue by stream and quarter (actual + forecast) for Tableau"),
    ("Tableau_Guidance", "Long-format guidance vs. actual/forecast for Tableau"),
]
for i, (t, d) in enumerate(tabs):
    put(ws, f"A{8+i}", t, f_bold); put(ws, f"B{8+i}", d)

r = 8 + len(tabs) + 1
section(ws, r, "Color legend", 2)
put(ws, f"A{r+1}", 1234, f_in, NUM); put(ws, f"B{r+1}", "Blue = hardcoded input from a filing or an analyst assumption (edit these)")
put(ws, f"A{r+2}", 1234, f_base, NUM); put(ws, f"B{r+2}", "Black = formula")
put(ws, f"A{r+3}", 1234, f_link, NUM); put(ws, f"B{r+3}", "Green = link to another tab")
put(ws, f"A{r+4}", "", fill=fill_key); put(ws, f"B{r+4}", "Yellow fill = key lever (scenario selector)")
put(ws, f"A{r+5}", "", fill=fill_fc); put(ws, f"B{r+5}", "Grey fill = forecast period")

r += 7
section(ws, r, "Disclaimer", 2)
put(ws, f"B{r+1}", "Independent portfolio project built solely from public information (Qualcomm earnings releases and SEC filings). "
    "Not affiliated with or endorsed by Qualcomm. Forecast drivers are analyst assumptions, not company data. Not investment advice.")
ws[f"B{r+1}"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[r + 1].height = 42

# =====================================================================
# Assumptions
# =====================================================================
ws = wb.create_sheet("Assumptions")
setup(ws, "Assumptions", "All forecast drivers. Change B5 to switch scenarios; edit blue cells to test your own view.",
      {"B": 14, "C": 11, "D": 11, "E": 11, "F": 11, "G": 100})
put(ws, "A5", "Scenario selector (1=Base, 2=Bull, 3=Bear)", f_bold)
put(ws, "B5", 1, f_in, fill=fill_key, align=CENTER)
put(ws, "C5", '=CHOOSE(B5,"Base","Bull","Bear")', f_bold)
ws["B5"].comment = Comment("Enter 1, 2 or 3. Every forecast, summary and Tableau export tab follows this cell.", "Model")

hdr = ["Driver", "Applies to", "Base", "Bull", "Bear", "Active", "Rationale / source"]
for j, t in enumerate(hdr):
    put(ws, f"{L(1+j)}7", t, f_hdr, fill=fill_hdr, align=CENTER if j else None)

DRV = [  # key, label, period, base, bull, bear, rationale
    ("h26", "Handsets revenue, YoY growth", "FY26 Q4", -0.27, -0.22, -0.32,
     "Calibrated so Base QCT lands near the FY26 Q4 guidance midpoint ($8.4B-$9.0B, S11). Handsets fell 13% (Q2) and 20% (Q3) YoY; "
     "management expects iPhone modem share materially below its prior 20% estimate, and FY25 Q4 is a tough comparison."),
    ("h27", "Handsets revenue, YoY growth", "FY27", -0.12, -0.05, -0.20,
     "Analyst assumption: continued Apple modem roll-off and memory-cost pressure on Android OEM builds, easing as comparisons reset."),
    ("a26", "Automotive revenue, YoY growth", "FY26 Q4", 0.60, 0.65, 0.50,
     "Management commentary after Q3 FY26 pointed to roughly 60% YoY automotive growth in Q4; FY26 Q3 grew 61% (S11)."),
    ("a27", "Automotive revenue, YoY growth", "FY27", 0.30, 0.40, 0.20,
     "Analyst assumption: growth decelerates off a larger base; exit-rate target raised to ~$7B annualized for FY26."),
    ("i26", "IoT revenue, YoY growth", "FY26 Q4", 0.09, 0.12, 0.05,
     "IoT grew 9% YoY in both FY26 Q2 and Q3 (S10, S11)."),
    ("i27", "IoT revenue, YoY growth", "FY27", 0.10, 0.15, 0.05,
     "Analyst assumption: industrial/robotics and edge-AI product cycles; Bull adds early data center contribution."),
    ("t26", "QTL revenue, YoY growth", "FY26 Q4", -0.08, -0.01, -0.14,
     "Calibrated to FY26 Q4 QTL guidance $1.2B-$1.4B (S11) vs. $1,409M in FY25 Q4."),
    ("t27", "QTL revenue, YoY growth", "FY27", -0.03, 0.01, -0.07,
     "Analyst assumption: licensing tracks handset units; FY26 Q3 QTL was down 3% YoY."),
    ("qm26", "QCT EBT margin", "FY26 Q4", 0.24, 0.25, 0.23,
     "FY26 Q4 guidance: QCT EBT margin 23%-25% (S11)."),
    ("qm27", "QCT EBT margin", "FY27", 0.26, 0.29, 0.23,
     "Analyst assumption: FY24-FY26 quarterly QCT margins ranged 26%-32%; Bear keeps memory-cost pressure."),
    ("lm26", "QTL EBT margin", "FY26 Q4", 0.70, 0.72, 0.67,
     "QTL EBT margin ran 69%-77% FY24-FY26 (Actuals tab)."),
    ("lm27", "QTL EBT margin", "FY27", 0.70, 0.73, 0.67,
     "Analyst assumption: stable licensing margin."),
]
AREF = {}
for i, (k, labl, per, b, u, d, why) in enumerate(DRV):
    r = 8 + i
    put(ws, f"A{r}", labl); put(ws, f"B{r}", per, align=CENTER)
    for cl, v in zip("CDE", (b, u, d)):
        put(ws, f"{cl}{r}", v, f_in, PCT, align=CENTER)
    put(ws, f"F{r}", f"=CHOOSE($B$5,C{r},D{r},E{r})", f_bold, PCT, align=CENTER)
    put(ws, f"G{r}", why)
    AREF[k] = r
put(ws, f"A{9+len(DRV)}", "Method: each forecast quarter = same quarter prior year x (1 + YoY growth). Using YoY on the same quarter "
    "preserves Qualcomm's seasonality (strong fiscal Q1 holiday handset builds).", f_sub)
ws.freeze_panes = "A8"

def AS(k, c="F"):
    return f"Assumptions!${c}${AREF[k]}"

# =====================================================================
# Actuals
# =====================================================================
ws = wb.create_sheet("Actuals")
setup(ws, "Actuals", "Reported results, $ in millions. Source: Qualcomm quarterly earnings releases (see docs/sources.csv).")
qwidths(ws, N_ACT)
header_row(ws, 4, FC0, [lab(*Q[i]) for i in range(N_ACT)])
put(ws, "A5", "Period end"); put(ws, "A6", "Source ID")
for i in range(N_ACT):
    k = Q[i]; c = col(i)
    put(ws, f"{c}5", A[k]["pe"], align=CENTER); put(ws, f"{c}6", A[k]["src"], align=CENTER)

last = FC0 + N_ACT - 1
section(ws, 8, "Revenue ($mm)", last)
rows = [(9, "Handsets", ("QCT", "Handsets", "rev")), (10, "Automotive", ("QCT", "Automotive", "rev")),
        (11, "IoT", ("QCT", "IoT", "rev")), (12, "QCT revenue (reported)", ("QCT", "Total", "rev")),
        (13, "QTL revenue", ("QTL", "Total", "rev"))]
for r, name, key in rows:
    put(ws, f"A{r}", name, f_bold if r == 12 else f_base); put(ws, f"B{r}", "$mm", align=CENTER)
    for i in range(N_ACT):
        put(ws, f"{col(i)}{r}", A[Q[i]][key], f_in, NUM)
put(ws, "A14", "QCT + QTL revenue", f_bold); put(ws, "B14", "$mm", align=CENTER)
for i in range(N_ACT):
    c = col(i); put(ws, f"{c}14", f"={c}12+{c}13", f_bold, NUM, border=top_border)

section(ws, 16, "Earnings before taxes ($mm)", last)
for r, name, key in [(17, "QCT EBT", ("QCT", "Total", "ebt")), (18, "QTL EBT", ("QTL", "Total", "ebt"))]:
    put(ws, f"A{r}", name); put(ws, f"B{r}", "$mm", align=CENTER)
    for i in range(N_ACT):
        put(ws, f"{col(i)}{r}", A[Q[i]][key], f_in, NUM)
put(ws, "A19", "QCT + QTL EBT", f_bold); put(ws, "B19", "$mm", align=CENTER)
for i in range(N_ACT):
    c = col(i); put(ws, f"{c}19", f"={c}17+{c}18", f_bold, NUM, border=top_border)

section(ws, 21, "Margins", last)
for r, name, num, den in [(22, "QCT EBT margin", 17, 12), (23, "QTL EBT margin", 18, 13)]:
    put(ws, f"A{r}", name); put(ws, f"B{r}", "%", align=CENTER)
    for i in range(N_ACT):
        c = col(i); put(ws, f"{c}{r}", f"=IF({c}{den}=0,0,{c}{num}/{c}{den})", fmt=PCT)

section(ws, 25, "YoY growth (same quarter prior year)", last)
for r, name, src_r in [(26, "Handsets", 9), (27, "Automotive", 10), (28, "IoT", 11), (29, "QCT revenue", 12), (30, "QTL revenue", 13)]:
    put(ws, f"A{r}", name); put(ws, f"B{r}", "%", align=CENTER)
    for i in range(N_ACT):
        c = col(i)
        if i < 4:
            put(ws, f"{c}{r}", "n/a", align=CENTER)
        else:
            p = col(i - 4)
            put(ws, f"{c}{r}", f"=IF({p}{src_r}=0,0,{c}{src_r}/{p}{src_r}-1)", fmt=PCT)

section(ws, 32, "QCT revenue mix", last)
for r, name, f in [(33, "Handsets % of QCT", "={c}9/{c}12"), (34, "Automotive + IoT % of QCT", "=({c}10+{c}11)/{c}12")]:
    put(ws, f"A{r}", name); put(ws, f"B{r}", "%", align=CENTER)
    for i in range(N_ACT):
        c = col(i); put(ws, f"{c}{r}", f.format(c=c), fmt=PCT)
ws.freeze_panes = "C5"

# =====================================================================
# Guidance
# =====================================================================
ws = wb.create_sheet("Guidance")
setup(ws, "Guidance", "Revenue guidance issued in the prior quarter's earnings release, $ in millions. "
      "FY24 Q1 guidance not collected (see docs/assumptions.md).")
NG = 12  # FY24Q1..FY26Q4 columns
qwidths(ws, NG)
header_row(ws, 4, FC0, [lab(*Q[i]) for i in range(NG)])
put(ws, "A5", "Issued (release date)"); put(ws, "A6", "Source ID")
blocks = [(8, "Total revenue (GAAP)", "Total_revenue"), (13, "QCT revenue", "QCT_revenue"), (18, "QTL revenue", "QTL_revenue")]
GROW = {}
for i in range(NG):
    k = Q[i]; c = col(i)
    if k in G:
        put(ws, f"{c}5", G[k]["issued"], align=CENTER); put(ws, f"{c}6", G[k]["src"], align=CENTER)
    else:
        put(ws, f"{c}5", "n/a", align=CENTER); put(ws, f"{c}6", "n/a", align=CENTER)
for r0, name, m in blocks:
    section(ws, r0, name, FC0 + NG - 1)
    for j, part in enumerate(("Low", "High", "Midpoint")):
        r = r0 + 1 + j
        put(ws, f"A{r}", f"  {part}", f_bold if part == "Midpoint" else f_base); put(ws, f"B{r}", "$mm", align=CENTER)
        for i in range(NG):
            k = Q[i]; c = col(i)
            if k not in G:
                put(ws, f"{c}{r}", "n/a", align=CENTER); continue
            if part == "Midpoint":
                put(ws, f"{c}{r}", f"=AVERAGE({c}{r0+1}:{c}{r0+2})", f_bold, NUM)
            else:
                put(ws, f"{c}{r}", G[k][m][j], f_in, NUM)
    GROW[m] = (r0 + 1, r0 + 2, r0 + 3)
put(ws, "A23", "Note: FY25 Q2 GAAP guidance includes dispute-settlement revenue; Non-GAAP range was $10.2B-$11.0B (S5).", f_sub)
ws.freeze_panes = "C5"

# =====================================================================
# Variance
# =====================================================================
ws = wb.create_sheet("Variance")
setup(ws, "Guidance vs. Actual Variance", "Actual revenue vs. guidance midpoint, $ in millions. Positive = beat.")
qwidths(ws, N_ACT)
VI = list(range(1, N_ACT))  # FY24Q2..FY26Q3
header_row(ws, 4, FC0, [lab(*Q[i]) for i in range(N_ACT)])
put(ws, f"{col(0)}5", "n/a", align=CENTER)
VROW = {}
for r0, name, m, arow in [(6, "QCT", "QCT_revenue", 12), (16, "QTL", "QTL_revenue", 13)]:
    lo, hi, mid = GROW[m]
    section(ws, r0, f"{name} revenue", FC0 + N_ACT - 1)
    labels = ["Actual", "Guidance low", "Guidance high", "Guidance midpoint", "Variance vs. midpoint ($mm)",
              "Variance vs. midpoint (%)", "Position in range", "Position (0% = low end, 100% = high end)"]
    for j, t in enumerate(labels):
        r = r0 + 1 + j
        put(ws, f"A{r}", t, f_bold if j in (4, 5) else f_base)
        put(ws, f"B{r}", "%" if j in (5, 7) else ("" if j == 6 else "$mm"), align=CENTER)
        for i in VI:
            c = col(i)
            f = [f"=Actuals!{c}{arow}", f"=Guidance!{c}{lo}", f"=Guidance!{c}{hi}", f"=Guidance!{c}{mid}",
                 f"={c}{r0+1}-{c}{r0+4}", f"=IF({c}{r0+4}=0,0,{c}{r0+5}/{c}{r0+4})",
                 f'=IF({c}{r0+1}>{c}{r0+3},"Above",IF({c}{r0+1}<{c}{r0+2},"Below","Within"))',
                 f"=IF({c}{r0+3}={c}{r0+2},0,({c}{r0+1}-{c}{r0+2})/({c}{r0+3}-{c}{r0+2}))"][j]
            fnt = f_link if j < 4 else (f_bold if j in (4, 5) else f_base)
            put(ws, f"{c}{r}", f, fnt, PCT if j in (5, 7) else (None if j == 6 else NUM),
                align=CENTER if j == 6 else None)
    VROW[name] = r0

section(ws, 26, "Scorecard (FY24 Q2 - FY26 Q3)", FC0 + N_ACT - 1)
c1, c2 = col(VI[0]), col(VI[-1])
put(ws, "A27", "Metric", f_hdr, fill=fill_hdr); put(ws, "B27", "QCT", f_hdr, fill=fill_hdr, align=CENTER)
put(ws, "C27", "QTL", f_hdr, fill=fill_hdr, align=CENTER)
score = [("Quarters above guidance range", 'COUNTIF({c1}{r}:{c2}{r},"Above")', 7, None),
         ("Quarters within guidance range", 'COUNTIF({c1}{r}:{c2}{r},"Within")', 7, None),
         ("Quarters below guidance range", 'COUNTIF({c1}{r}:{c2}{r},"Below")', 7, None),
         ("Average variance vs. midpoint (%)", "AVERAGE({c1}{r}:{c2}{r})", 6, PCT),
         ("Average position in range", "AVERAGE({c1}{r}:{c2}{r})", 8, PCT),
         ("Largest beat vs. midpoint ($mm)", "MAX({c1}{r}:{c2}{r})", 5, NUM),
         ("Largest miss vs. midpoint ($mm)", "MIN({c1}{r}:{c2}{r})", 5, NUM)]
for j, (t, f, off, fmt) in enumerate(score):
    r = 28 + j
    put(ws, f"A{r}", t)
    for cl, seg in (("B", "QCT"), ("C", "QTL")):
        put(ws, f"{cl}{r}", "=" + f.format(c1=c1, c2=c2, r=VROW[seg] + off), f_bold, fmt, align=CENTER)
ws.freeze_panes = "C5"

# =====================================================================
# Forecast
# =====================================================================
ws = wb.create_sheet("Forecast")
setup(ws, "Quarterly Forecast", "Actuals FY24 Q1 - FY26 Q3 (linked) + forecast FY26 Q4 - FY27 Q4 (grey). $ in millions.")
NQ = len(Q); qwidths(ws, NQ)
header_row(ws, 4, FC0, [lab(*q) for q in Q])
put(ws, "A5", "Type"); put(ws, "A6", "Scenario")
LASTC = FC0 + NQ - 1


def g_key(i, base):  # pick FY26Q4 vs FY27 driver
    return base + ("26" if Q[i] == (2026, 4) else "27")


for i in range(NQ):
    c = col(i); fc = i >= N_ACT
    put(ws, f"{c}5", "Forecast" if fc else "Actual", f_bold if fc else f_base, align=CENTER)
    put(ws, f"{c}6", "=Assumptions!$C$5" if fc else "Reported", f_link if fc else f_base, align=CENTER)

section(ws, 8, "Revenue ($mm)", LASTC)
FR = {"h": 9, "a": 10, "i": 11, "qct": 12, "qtl": 13, "tot": 14, "qe": 17, "le": 18, "te": 19,
      "qm": 22, "lm": 23, "gy": 26, "nh": 27, "nhs": 28, "nhg": 29, "hg": 30}
for key, name, arow, drv in [("h", "Handsets", 9, "h"), ("a", "Automotive", 10, "a"), ("i", "IoT", 11, "i")]:
    r = FR[key]; put(ws, f"A{r}", name); put(ws, f"B{r}", "$mm", align=CENTER)
    for i in range(NQ):
        c = col(i)
        if i < N_ACT:
            put(ws, f"{c}{r}", f"=Actuals!{c}{arow}", f_link, NUM)
        else:
            put(ws, f"{c}{r}", f"={col(i-4)}{r}*(1+{AS(g_key(i, drv))})", fmt=NUM, fill=fill_fc)
r = FR["qct"]; put(ws, f"A{r}", "QCT revenue", f_bold); put(ws, f"B{r}", "$mm", align=CENTER)
for i in range(NQ):
    c = col(i)
    if i < N_ACT:
        put(ws, f"{c}{r}", f"=Actuals!{c}12", Font(name=F, size=10, color="008000", bold=True), NUM, border=top_border)
    else:
        put(ws, f"{c}{r}", f"=SUM({c}9:{c}11)", f_bold, NUM, fill=fill_fc, border=top_border)
r = FR["qtl"]; put(ws, f"A{r}", "QTL revenue"); put(ws, f"B{r}", "$mm", align=CENTER)
for i in range(NQ):
    c = col(i)
    if i < N_ACT:
        put(ws, f"{c}{r}", f"=Actuals!{c}13", f_link, NUM)
    else:
        put(ws, f"{c}{r}", f"={col(i-4)}{r}*(1+{AS(g_key(i, 't'))})", fmt=NUM, fill=fill_fc)
r = FR["tot"]; put(ws, f"A{r}", "QCT + QTL revenue", f_bold); put(ws, f"B{r}", "$mm", align=CENTER)
for i in range(NQ):
    c = col(i); put(ws, f"{c}{r}", f"={c}12+{c}13", f_bold, NUM, fill=fill_fc if i >= N_ACT else None, border=top_border)

section(ws, 16, "Earnings before taxes ($mm)", LASTC)
for key, name, arow, rev, drv in [("qe", "QCT EBT", 17, 12, "qm"), ("le", "QTL EBT", 18, 13, "lm")]:
    r = FR[key]; put(ws, f"A{r}", name); put(ws, f"B{r}", "$mm", align=CENTER)
    for i in range(NQ):
        c = col(i)
        if i < N_ACT:
            put(ws, f"{c}{r}", f"=Actuals!{c}{arow}", f_link, NUM)
        else:
            put(ws, f"{c}{r}", f"={c}{rev}*{AS(g_key(i, drv))}", fmt=NUM, fill=fill_fc)
r = FR["te"]; put(ws, f"A{r}", "QCT + QTL EBT", f_bold); put(ws, f"B{r}", "$mm", align=CENTER)
for i in range(NQ):
    c = col(i); put(ws, f"{c}{r}", f"={c}17+{c}18", f_bold, NUM, fill=fill_fc if i >= N_ACT else None, border=top_border)

section(ws, 21, "Margins", LASTC)
for key, name, num, den in [("qm", "QCT EBT margin", 17, 12), ("lm", "QTL EBT margin", 18, 13)]:
    r = FR[key]; put(ws, f"A{r}", name); put(ws, f"B{r}", "%", align=CENTER)
    for i in range(NQ):
        c = col(i); put(ws, f"{c}{r}", f"=IF({c}{den}=0,0,{c}{num}/{c}{den})", fmt=PCT, fill=fill_fc if i >= N_ACT else None)

section(ws, 25, "Diversification metrics", LASTC)
put(ws, "A26", "QCT + QTL revenue, YoY growth"); put(ws, "B26", "%", align=CENTER)
put(ws, "A27", "Non-handset QCT revenue (Auto + IoT)", f_bold); put(ws, "B27", "$mm", align=CENTER)
put(ws, "A28", "Non-handset share of QCT"); put(ws, "B28", "%", align=CENTER)
put(ws, "A29", "Non-handset QCT revenue, YoY growth"); put(ws, "B29", "%", align=CENTER)
put(ws, "A30", "Handsets, YoY growth"); put(ws, "B30", "%", align=CENTER)
for i in range(NQ):
    c = col(i); fl = fill_fc if i >= N_ACT else None
    put(ws, f"{c}27", f"={c}10+{c}11", f_bold, NUM, fill=fl)
    put(ws, f"{c}28", f"=IF({c}12=0,0,{c}27/{c}12)", fmt=PCT, fill=fl)
    if i < 4:
        for rr in (26, 29, 30):
            put(ws, f"{c}{rr}", "n/a", align=CENTER)
    else:
        p = col(i - 4)
        put(ws, f"{c}26", f"=IF({p}14=0,0,{c}14/{p}14-1)", fmt=PCT, fill=fl)
        put(ws, f"{c}29", f"=IF({p}27=0,0,{c}27/{p}27-1)", fmt=PCT, fill=fl)
        put(ws, f"{c}30", f"=IF({p}9=0,0,{c}9/{p}9-1)", fmt=PCT, fill=fl)
put(ws, "A32", "Method: forecast quarter = same quarter prior year x (1 + YoY growth from Assumptions). "
    "QCT forecast = sum of streams; EBT = revenue x EBT margin.", f_sub)
ws.freeze_panes = "C5"

# =====================================================================
# Summary (annual)
# =====================================================================
ws = wb.create_sheet("Summary")
setup(ws, "Fiscal Year Summary", "Fiscal years end late September. FY26E = 3 reported quarters + Q4 forecast. $ in millions.",
      {"C": 13, "D": 13, "E": 13, "F": 13})
yrs = [("FY24", 0, "Actual"), ("FY25", 4, "Actual"), ("FY26E", 8, "Act + Fcst"), ("FY27E", 12, "Forecast")]
put(ws, "A4", "Line item", f_hdr, fill=fill_hdr); put(ws, "B4", "Unit", f_hdr, fill=fill_hdr, align=CENTER)
for j, (y, s, t) in enumerate(yrs):
    put(ws, f"{L(3+j)}4", y, f_hdr, fill=fill_hdr, align=CENTER)
    put(ws, f"{L(3+j)}5", t, f_sub, align=CENTER)
put(ws, "A5", "Basis")
put(ws, "A6", "Scenario"); put(ws, "C6", "=Assumptions!$C$5", f_link)

SR = [(8, "Handsets", 9), (9, "Automotive", 10), (10, "IoT", 11), (11, "QCT revenue", 12), (12, "QTL revenue", 13),
      (13, "QCT + QTL revenue", 14), (15, "QCT EBT", 17), (16, "QTL EBT", 18), (17, "QCT + QTL EBT", 19)]
section(ws, 7, "Revenue and EBT ($mm)", 6)
for r, name, frow in SR:
    bold = name.startswith("QCT +") or name == "QCT revenue"
    put(ws, f"A{r}", name, f_bold if bold else f_base); put(ws, f"B{r}", "$mm", align=CENTER)
    for j, (y, s, t) in enumerate(yrs):
        put(ws, f"{L(3+j)}{r}", f"=SUM(Forecast!{col(s)}{frow}:{col(s+3)}{frow})", f_bold if bold else f_base, NUM,
            border=top_border if name.startswith("QCT +") else None)
put(ws, "A14", "")
section(ws, 19, "Growth and margins", 6)
gro = [(20, "Handsets, YoY", 8), (21, "Automotive, YoY", 9), (22, "IoT, YoY", 10), (23, "QCT revenue, YoY", 11),
       (24, "QTL revenue, YoY", 12), (25, "QCT + QTL revenue, YoY", 13)]
for r, name, sr in gro:
    put(ws, f"A{r}", name); put(ws, f"B{r}", "%", align=CENTER); put(ws, f"C{r}", "n/a", align=CENTER)
    for j in range(1, 4):
        c, p = L(3 + j), L(2 + j)
        put(ws, f"{c}{r}", f"=IF({p}{sr}=0,0,{c}{sr}/{p}{sr}-1)", fmt=PCT, align=CENTER)
for r, name, n, d in [(26, "QCT EBT margin", 15, 11), (27, "QTL EBT margin", 16, 12)]:
    put(ws, f"A{r}", name); put(ws, f"B{r}", "%", align=CENTER)
    for j in range(4):
        c = L(3 + j); put(ws, f"{c}{r}", f"=IF({c}{d}=0,0,{c}{n}/{c}{d})", fmt=PCT, align=CENTER)

section(ws, 29, "Core question: does diversification offset handset decline?", 6)
put(ws, "A30", "Change in handset revenue vs. prior year"); put(ws, "B30", "$mm", align=CENTER)
put(ws, "A31", "Change in Auto + IoT revenue vs. prior year"); put(ws, "B31", "$mm", align=CENTER)
put(ws, "A32", "Change in QTL revenue vs. prior year"); put(ws, "B32", "$mm", align=CENTER)
put(ws, "A33", "Net change in QCT + QTL revenue", f_bold); put(ws, "B33", "$mm", align=CENTER)
put(ws, "A34", "Offset ratio (Auto + IoT gain / handset loss)"); put(ws, "B34", "x", align=CENTER)
put(ws, "A35", "Non-handset share of QCT revenue"); put(ws, "B35", "%", align=CENTER)
put(ws, "A36", "Verdict", f_bold)
put(ws, "C30", "n/a", align=CENTER); put(ws, "C31", "n/a", align=CENTER); put(ws, "C32", "n/a", align=CENTER)
put(ws, "C33", "n/a", align=CENTER); put(ws, "C34", "n/a", align=CENTER); put(ws, "C36", "n/a", align=CENTER)
put(ws, "C35", "=IF(C11=0,0,(C9+C10)/C11)", fmt=PCT, align=CENTER)
for j in range(1, 4):
    c, p = L(3 + j), L(2 + j)
    put(ws, f"{c}30", f"={c}8-{p}8", fmt=NUM, align=CENTER)
    put(ws, f"{c}31", f"=({c}9+{c}10)-({p}9+{p}10)", fmt=NUM, align=CENTER)
    put(ws, f"{c}32", f"={c}12-{p}12", fmt=NUM, align=CENTER)
    put(ws, f"{c}33", f"=SUM({c}30:{c}32)", f_bold, NUM, align=CENTER, border=top_border)
    put(ws, f"{c}34", f'=IF({c}30>=0,"No decline",{c}31/-{c}30)', fmt='0.00"x"', align=CENTER)
    put(ws, f"{c}35", f"=IF({c}11=0,0,({c}9+{c}10)/{c}11)", fmt=PCT, align=CENTER)
    put(ws, f"{c}36", f'=IF({c}30>=0,"Handsets grew",IF({c}31>=-{c}30,"Offset","Not offset"))', f_bold, align=CENTER)
put(ws, "A38", "Offset ratio >= 1.00x means Automotive + IoT added at least as many dollars as Handsets lost that year.", f_sub)
ws.freeze_panes = "C5"

# =====================================================================
# Scenarios (side by side, independent of selector)
# =====================================================================
ws = wb.create_sheet("Scenarios")
setup(ws, "Scenario Comparison", "Computed directly from each scenario's drivers, so all three show at once regardless of the selector. $ in millions.",
      {"C": 13, "D": 13, "E": 13})
put(ws, "A4", "Line item", f_hdr, fill=fill_hdr); put(ws, "B4", "Unit", f_hdr, fill=fill_hdr, align=CENTER)
for j, s in enumerate(("Base", "Bull", "Bear")):
    put(ws, f"{L(3+j)}4", s, f_hdr, fill=fill_hdr, align=CENTER)
# FY26 Q1-Q3 actual sums (Forecast cols for i=8..10) and FY25 Q4 (i=7)
f26 = lambda frow: f"SUM(Forecast!{col(8)}{frow}:{col(10)}{frow})"
f25q4 = lambda frow: f"Forecast!{col(7)}{frow}"
segs = [("Handsets", 9, "h"), ("Automotive", 10, "a"), ("IoT", 11, "i"), ("QTL revenue", 13, "t")]
section(ws, 5, "FY26E revenue ($mm)", 5)
section(ws, 12, "FY27E revenue ($mm)", 5)
for j, sc in enumerate("CDE"):
    for k, (name, frow, d) in enumerate(segs):
        r26, r27 = 6 + k, 13 + k
        q4 = f"{f25q4(frow)}*(1+Assumptions!${sc}${AREF[d+'26']})"
        put(ws, f"A{r26}", name); put(ws, f"B{r26}", "$mm", align=CENTER)
        put(ws, f"{sc}{r26}", f"={f26(frow)}+{q4}", fmt=NUM)
        put(ws, f"A{r27}", name); put(ws, f"B{r27}", "$mm", align=CENTER)
        put(ws, f"{sc}{r27}", f"=({f26(frow)}+{q4})*(1+Assumptions!${sc}${AREF[d+'27']})", fmt=NUM)
    put(ws, "A10", "QCT + QTL revenue", f_bold); put(ws, "B10", "$mm", align=CENTER)
    put(ws, f"{sc}10", f"=SUM({sc}6:{sc}9)", f_bold, NUM, border=top_border)
    put(ws, "A17", "QCT + QTL revenue", f_bold); put(ws, "B17", "$mm", align=CENTER)
    put(ws, f"{sc}17", f"=SUM({sc}13:{sc}16)", f_bold, NUM, border=top_border)
section(ws, 19, "FY27E profitability and growth", 5)
rows27 = [(20, "QCT revenue", "$mm", "={c}13+{c}14+{c}15", NUM),
          (21, "QCT EBT", "$mm", f"={{c}}20*Assumptions!${{c}}${AREF['qm27']}", NUM),
          (22, "QTL EBT", "$mm", f"={{c}}16*Assumptions!${{c}}${AREF['lm27']}", NUM),
          (23, "QCT + QTL EBT", "$mm", "={c}21+{c}22", NUM),
          (24, "QCT + QTL revenue, YoY vs. FY26E", "%", "=IF({c}10=0,0,{c}17/{c}10-1)", PCT),
          (25, "Non-handset share of QCT, FY27E", "%", "=IF({c}20=0,0,({c}14+{c}15)/{c}20)", PCT),
          (26, "Handsets $ change, FY27E vs. FY26E", "$mm", "={c}13-{c}6", NUM),
          (27, "Auto + IoT $ change, FY27E vs. FY26E", "$mm", "=({c}14+{c}15)-({c}7+{c}8)", NUM),
          (28, "Verdict", "", '=IF({c}26>=0,"Handsets grew",IF({c}27>=-{c}26,"Offset","Not offset"))', None)]
for r, name, u, f, fmt in rows27:
    put(ws, f"A{r}", name, f_bold if name.startswith(("QCT +", "Verdict")) else f_base)
    put(ws, f"B{r}", u, align=CENTER)
    for sc in "CDE":
        put(ws, f"{sc}{r}", f.format(c=sc), f_bold if name.startswith(("QCT +", "Verdict")) else f_base, fmt, align=CENTER)
put(ws, "A30", "Tie-out: the column matching the active scenario equals Summary FY26E / FY27E (see Checks).", f_sub)

# =====================================================================
# Sensitivity
# =====================================================================
ws = wb.create_sheet("Sensitivity")
setup(ws, "Sensitivity: FY27E Revenue", "QCT + QTL revenue ($mm) for FY27E by handset growth (rows) x automotive growth (columns). "
      "IoT and QTL use the active scenario.", {"B": 14})
put(ws, "A4", "FY26E bases (active scenario)", f_bold)
bases = [(5, "Handsets", "C8"), (6, "Automotive", "C9"), (7, "IoT", "C10"), (8, "QTL revenue", "C12")]
for r, name, ref in bases:
    put(ws, f"A{r}", name); put(ws, f"B{r}", f"=Summary!E{ref[1:]}", f_link, NUM)
put(ws, "A9", "IoT FY27 growth (active)"); put(ws, "B9", f"={AS('i27')}", f_link, PCT)
put(ws, "A10", "QTL FY27 growth (active)"); put(ws, "B10", f"={AS('t27')}", f_link, PCT)
put(ws, "A12", "Handsets YoY  \\  Automotive YoY", f_hdr, fill=fill_hdr)
put(ws, "B12", "", fill=fill_hdr)
auto_g = [0.10, 0.20, 0.30, 0.40, 0.50]
hand_g = [-0.25, -0.20, -0.15, -0.10, -0.05, 0.0, 0.05]
for j, g in enumerate(auto_g):
    put(ws, f"{L(3+j)}12", g, Font(name=F, size=10, bold=True, color="0000FF"), PCT, fill=fill_sec, align=CENTER)
    ws.column_dimensions[L(3 + j)].width = 13
for i, g in enumerate(hand_g):
    r = 13 + i
    put(ws, f"B{r}", g, Font(name=F, size=10, bold=True, color="0000FF"), PCT, fill=fill_sec, align=CENTER)
    for j in range(len(auto_g)):
        c = L(3 + j)
        put(ws, f"{c}{r}", f"=$B$5*(1+$B{r})+$B$6*(1+{c}$12)+$B$7*(1+$B$9)+$B$8*(1+$B$10)", fmt=NUM, align=CENTER)
put(ws, "A21", "FY26E QCT + QTL revenue (for reference)"); put(ws, "B21", "=Summary!E13", f_link, NUM)
put(ws, "A22", "Shade guide: compare each cell with B21 to see which combinations grow revenue in FY27.", f_sub)

# =====================================================================
# Checks
# =====================================================================
ws = wb.create_sheet("Checks")
setup(ws, "Model Checks", "Every check should read OK.", {"B": 70, "C": 12})
put(ws, "A4", "OVERALL", f_bold)
ca, cz = col(0), col(N_ACT - 1)
checks = [
    ("QCT streams sum to reported QCT revenue (all actual quarters)",
     f'=IF(SUMPRODUCT(ABS(Actuals!{ca}9:{cz}9+Actuals!{ca}10:{cz}10+Actuals!{ca}11:{cz}11-Actuals!{ca}12:{cz}12))<1,"OK","ERROR")'),
    ("Guidance midpoints equal (low + high) / 2",
     "=IF(" + "+".join(f"SUMPRODUCT(ABS(Guidance!{col(1)}{m}:{col(11)}{m}-(Guidance!{col(1)}{lo}:{col(11)}{lo}+Guidance!{col(1)}{hi}:{col(11)}{hi})/2))"
                       for lo, hi, m in GROW.values()) + '<1,"OK","ERROR")'),
    ("Guidance low <= high for every quarter",
     "=IF(" + "+".join(f"SUMPRODUCT(--(Guidance!{col(1)}{lo}:{col(11)}{lo}>Guidance!{col(1)}{hi}:{col(11)}{hi}))"
                       for lo, hi, m in GROW.values()) + '=0,"OK","ERROR")'),
    ("Scenario selector is 1, 2 or 3",
     '=IF(OR(Assumptions!B5=1,Assumptions!B5=2,Assumptions!B5=3),"OK","ERROR")'),
    ("Summary annual revenue ties to quarterly Forecast",
     f'=IF(ABS(SUM(Summary!C13:F13)-SUM(Forecast!{col(0)}14:{col(15)}14))<1,"OK","ERROR")'),
    ("Active-scenario column in Scenarios ties to Summary FY27E",
     '=IF(ABS(INDEX(Scenarios!C17:E17,Assumptions!B5)-Summary!F13)<1,"OK","ERROR")'),
    ("Forecast QCT = sum of forecast streams",
     f'=IF(ABS(SUM(Forecast!{col(N_ACT)}12:{col(15)}12)-SUM(Forecast!{col(N_ACT)}9:{col(15)}11))<1,"OK","ERROR")'),
]
put(ws, "A6", "#", f_hdr, fill=fill_hdr); put(ws, "B6", "Check", f_hdr, fill=fill_hdr); put(ws, "C6", "Result", f_hdr, fill=fill_hdr, align=CENTER)
for i, (t, f) in enumerate(checks):
    r = 7 + i
    put(ws, f"A{r}", i + 1, align=CENTER); put(ws, f"B{r}", t); put(ws, f"C{r}", f, f_bold, align=CENTER)
last_chk = 7 + len(checks) - 1
put(ws, "C4", f'=IF(COUNTIF(C7:C{last_chk},"OK")={len(checks)},"ALL OK","CHECK ERRORS")', f_bold, align=CENTER)
r = last_chk + 2
section(ws, r, "Informational", 3)
fq = col(N_ACT)  # FY26Q4 column
put(ws, f"B{r+1}", "FY26 Q4 forecast QCT revenue vs. company guidance range")
put(ws, f"C{r+1}", f'=IF(Forecast!{fq}12>Guidance!{fq}{GROW["QCT_revenue"][1]},"Above",IF(Forecast!{fq}12<Guidance!{fq}{GROW["QCT_revenue"][0]},"Below","Within"))',
    f_bold, align=CENTER)
put(ws, f"B{r+2}", "FY26 Q4 forecast QTL revenue vs. company guidance range")
put(ws, f"C{r+2}", f'=IF(Forecast!{fq}13>Guidance!{fq}{GROW["QTL_revenue"][1]},"Above",IF(Forecast!{fq}13<Guidance!{fq}{GROW["QTL_revenue"][0]},"Below","Within"))',
    f_bold, align=CENTER)

# =====================================================================
# Tableau exports (long format, formula-linked)
# =====================================================================
ws = wb.create_sheet("Tableau_Revenue")
cols = ["period", "fiscal_year", "fiscal_quarter", "period_sort", "type", "scenario", "segment", "stream", "revenue_musd", "stream_ebt_musd"]
for j, t in enumerate(cols):
    put(ws, f"{L(1+j)}1", t, f_hdr, fill=fill_hdr); ws.column_dimensions[L(1 + j)].width = 15
r = 2
streams = [("QCT", "Handsets", 9, None), ("QCT", "Automotive", 10, None), ("QCT", "IoT", 11, None), ("QTL", "Licensing", 13, 18)]
for i, (fy, q) in enumerate(Q):
    c = col(i)
    for seg, st, frow, erow in streams:
        vals = [lab(fy, q), fy, q, fy + q / 10, f"=Forecast!{c}5", f"=IF(E{r}=\"Actual\",\"Reported\",Assumptions!$C$5)",
                seg, st, f"=Forecast!{c}{frow}", f"=Forecast!{c}{erow}" if erow else ""]
        for j, v in enumerate(vals):
            put(ws, f"{L(1+j)}{r}", v, f_link if isinstance(v, str) and v.startswith("=") else f_base,
                NUM if j in (8, 9) else ("0.0" if j == 3 else None))
        r += 1
ws.freeze_panes = "A2"

ws = wb.create_sheet("Tableau_Guidance")
cols = ["period", "fiscal_year", "fiscal_quarter", "period_sort", "segment", "type", "value_musd", "guide_low_musd",
        "guide_high_musd", "guide_mid_musd", "variance_vs_mid_pct", "position"]
for j, t in enumerate(cols):
    put(ws, f"{L(1+j)}1", t, f_hdr, fill=fill_hdr); ws.column_dimensions[L(1 + j)].width = 16
r = 2
for i in range(1, 12):  # FY24Q2..FY26Q4
    fy, q = Q[i]; c = col(i)
    for seg, m, frow in (("QCT", "QCT_revenue", 12), ("QTL", "QTL_revenue", 13)):
        lo, hi, mid = GROW[m]
        vals = [lab(fy, q), fy, q, fy + q / 10, seg, f"=Forecast!{c}5", f"=Forecast!{c}{frow}",
                f"=Guidance!{c}{lo}", f"=Guidance!{c}{hi}", f"=Guidance!{c}{mid}",
                f"=IF(J{r}=0,0,G{r}/J{r}-1)", f'=IF(G{r}>I{r},"Above",IF(G{r}<H{r},"Below","Within"))']
        for j, v in enumerate(vals):
            fmt = NUM if j in (6, 7, 8, 9) else (PCT if j == 10 else ("0.0" if j == 3 else None))
            put(ws, f"{L(1+j)}{r}", v, f_link if isinstance(v, str) and v.startswith("=Forecast") or (isinstance(v, str) and v.startswith("=Guidance")) else f_base, fmt)
        r += 1
ws.freeze_panes = "A2"

wb.save(OUT)
print("saved", OUT)
